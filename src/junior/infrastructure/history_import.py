"""Import both Junior 1.x history workbook layouts into the 2.0 lifecycle DB."""

from __future__ import annotations

import sqlite3
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

OLD_HEADERS = (
    "History Type", "Company", "Role", "Source", "ATS Platform",
    "Work Arrangement", "Location", "Comp Range", "Event Date", "Status",
    "Outcome Category", "Recruiter/Contact", "Technical Match",
    "Hiring Probability", "Skills/Signals", "Primary Blocker",
    "Secondary Blocker", "Revisit", "Include In Job Radar", "Import Key", "Notes",
)
SIMPLE_HEADERS = (
    "Job Radar ID", "Date", "Company", "Role", "Posting URL", "Lead Source",
    "Decision", "Outcome", "Recruiter/Contact", "Notes", "Include In Job Radar",
)


@dataclass(frozen=True, slots=True)
class HistoryImportSummary:
    rows_read: int
    rows_imported: int
    rows_skipped: int


def import_history_workbook(
    database_path: str | Path, workbook_path: str | Path
) -> HistoryImportSummary:
    path = Path(workbook_path)
    if not path.is_file():
        raise ValueError(f"Job history workbook does not exist: {path}")
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheets = _matching_sheets(workbook)
    if not sheets:
        raise ValueError("Workbook does not match a Junior history schema.")
    read = imported = skipped = 0
    with sqlite3.connect(database_path) as connection:
        for sheet, headers, simplified in sheets:
            for row_number, row in enumerate(
                sheet.iter_rows(min_row=2, max_col=len(headers), values_only=True), 2
            ):
                values = dict(
                    zip(headers, (_text(value) for value in row), strict=True)
                )
                if all(value is None for value in values.values()):
                    continue
                read += 1
                if (values.get("Include In Job Radar") or "").casefold() == "no":
                    skipped += 1
                    continue
                record = (
                    _simple_record(values, sheet.title, row_number)
                    if simplified
                    else _old_record(values, sheet.title, row_number)
                )
                _upsert_history(connection, record)
                if simplified and record["history_type"] == "Pipeline":
                    _sync_application(connection, record)
                imported += 1
    return HistoryImportSummary(read, imported, skipped)


def _matching_sheets(workbook):
    matches = []
    for sheet in workbook.worksheets:
        headers = tuple(_text(cell.value) for cell in sheet[1])
        if headers[: len(SIMPLE_HEADERS)] == SIMPLE_HEADERS:
            matches.append((sheet, SIMPLE_HEADERS, True))
        elif (
            sheet.title in {"Pipeline Import", "Reviewed Import"}
            and headers[: len(OLD_HEADERS)] == OLD_HEADERS
        ):
            matches.append((sheet, OLD_HEADERS, False))
    return matches


def _old_record(values, sheet: str, row: int) -> dict[str, object]:
    for key in ("History Type", "Company", "Role", "Import Key"):
        if not values.get(key):
            raise ValueError(f"{sheet} row {row} is missing {key}.")
    return {
        "history_type": values["History Type"], "company": values["Company"],
        "role": values["Role"], "source": values["Source"],
        "ats_platform": values["ATS Platform"],
        "work_arrangement": values["Work Arrangement"], "location": values["Location"],
        "comp_range": values["Comp Range"], "event_date": values["Event Date"],
        "status": values["Status"], "outcome_category": values["Outcome Category"],
        "recruiter_contact": values["Recruiter/Contact"],
        "technical_match": values["Technical Match"],
        "hiring_probability": values["Hiring Probability"],
        "skills_signals": values["Skills/Signals"],
        "primary_blocker": values["Primary Blocker"],
        "secondary_blocker": values["Secondary Blocker"], "revisit": values["Revisit"],
        "import_key": values["Import Key"], "notes": values["Notes"],
        "posting_url": None, "job_radar_id": None,
    }


def _simple_record(values, sheet: str, row: int) -> dict[str, object]:
    company, role = values.get("Company"), values.get("Role")
    if not company or not role:
        raise ValueError(f"{sheet} row {row} is missing Company or Role.")
    decision, outcome = values.get("Decision"), values.get("Outcome")
    combined = f"{decision or ''} {outcome or ''}".casefold()
    history_type = (
        "Reviewed"
        if any(word in combined for word in ("skip", "avoid", "reviewed"))
        else "Pipeline"
    )
    job_id, url, event_date = (
        values.get("Job Radar ID"), values.get("Posting URL"), values.get("Date")
    )
    key = (
        f"job-radar-id:{_token(job_id)}" if job_id else
        f"posting-url:{_token(url)}" if url else
        (
            f"manual:{_token(company)}:{_token(role)}:"
            f"{_token(event_date or 'unknown')}:{row}"
        )
    )
    return {
        "history_type": history_type, "company": company, "role": role,
        "source": values.get("Lead Source"), "ats_platform": None,
        "work_arrangement": None, "location": None, "comp_range": None,
        "event_date": event_date, "status": decision, "outcome_category": outcome,
        "recruiter_contact": values.get("Recruiter/Contact"),
        "technical_match": None, "hiring_probability": None, "skills_signals": None,
        "primary_blocker": None, "secondary_blocker": None, "revisit": None,
        "import_key": key, "notes": values.get("Notes"), "posting_url": url,
        "job_radar_id": job_id,
    }


def _upsert_history(connection: sqlite3.Connection, record: dict[str, object]) -> None:
    columns = (
        "history_type", "company", "role", "source", "ats_platform",
        "work_arrangement", "location", "comp_range", "event_date", "status",
        "outcome_category", "recruiter_contact", "technical_match",
        "hiring_probability", "skills_signals", "primary_blocker",
        "secondary_blocker", "revisit", "import_key", "notes",
    )
    connection.execute(
        f"""INSERT INTO job_history ({', '.join(columns)})
            VALUES ({', '.join('?' for _ in columns)})
            ON CONFLICT(import_key) DO UPDATE SET
              history_type=excluded.history_type, company=excluded.company,
              role=excluded.role, status=excluded.status,
              outcome_category=excluded.outcome_category, notes=excluded.notes,
              updated_at=CURRENT_TIMESTAMP""",
        tuple(record[column] for column in columns),
    )


def _sync_application(
    connection: sqlite3.Connection, record: dict[str, object]
) -> None:
    identifier = str(record.get("job_radar_id") or record["import_key"])
    connection.execute(
        """INSERT INTO application_tracker (
               job_radar_id, company_name, role_title, source_url, status,
               outcome, notes, last_activity_on
           ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
           ON CONFLICT(job_radar_id) DO UPDATE SET
             status=excluded.status, outcome=excluded.outcome,
             notes=excluded.notes, last_activity_on=excluded.last_activity_on,
             updated_at=CURRENT_TIMESTAMP""",
        (
            identifier, record["company"], record["role"], record["posting_url"],
            record["status"] or "review_needed", record["outcome_category"],
            record["notes"], record["event_date"],
        ),
    )


def _text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = " ".join(str(value).split())
    return text or None


def _token(value: object) -> str:
    return "-".join(str(value).casefold().split())
